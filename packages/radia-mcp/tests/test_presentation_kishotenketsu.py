"""起承転結 check: the arc, the backup cut, and the inspired-by rule.

The module shipped without tests, and two of its assumptions turned out to be
wrong on the first real deck it was pointed at (IGTE'26, 2026-09-06):

  * `shapes.title` is None for a deck that draws its own title band, so every
    slide read as untitled and the "Backup" cut never fired -- nine hidden
    slides were counted into the arc and the turn was located on the summary.
  * the turn regex knew "use/add/replace" but not the imperative the lab
    actually writes ("Keep the ladder, add one surface mode, ...").

The inspired-by rule is Sugahara's, from the same deck: a turn that names the
outside idea it borrowed, with a credit, is followed on first hearing.
"""
import pytest

pptx = pytest.importorskip("pptx")

from pptx.util import Inches, Pt  # noqa: E402

from radia_mcp.presentation._kishotenketsu import (  # noqa: E402
    presentation_kishotenketsu_check,
)
from radia_mcp.presentation.tools import (  # noqa: E402
    presentation_check_qa_backup_slides,
)

# 起, 承, 転, 結, 結 -- the shape the check looks for, in the lab's own format:
# a short noun-phrase title, a body, and the claim in the bottom banner.
ARC = [
    ("The ladder", ["Kameari 2018 builds it."], "A circuit, exact at DC."),
    ("Its far end", ["Ten rungs."], "A finite ladder cannot make that slope."),
    ("One space", ["Inspired by XFEM [10]: basis and enrichment in one matrix."],
     "Keep the ladder, add one surface mode."),
    ("A cylinder", ["Copper, 5 mm."], "Two unknowns hold the band to 0.06 %."),
    ("Summary", ["Bulk plus surface."], "Keep the ladder, add the surface mode."),
]


def _deck(tmp_path, rows, hidden_from=None, titled=True, name="arc.pptx"):
    prs = pptx.Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    # slide 1 is the cover; the check drops it
    for index, (title, body, takeaway) in enumerate([("Cover", ["Talk"], "")] + rows):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if titled:
            head = slide.shapes.add_textbox(Inches(0.6), Inches(0.1),
                                            Inches(12.0), Inches(0.7))
            head.text_frame.text = title
        for line in body:
            box = slide.shapes.add_textbox(Inches(0.6), Inches(1.5),
                                           Inches(12.0), Inches(3.0))
            box.text_frame.text = line
        if takeaway:
            band = slide.shapes.add_textbox(Inches(0.4), Inches(6.4),
                                            Inches(12.5), Inches(0.6))
            band.text_frame.paragraphs[0].add_run().text = takeaway
        if hidden_from is not None and index + 1 >= hidden_from:
            slide._element.set("show", "0")
    out = tmp_path / name
    prs.save(str(out))
    return out


def test_the_four_phases_are_found(tmp_path):
    r = presentation_kishotenketsu_check(str(_deck(tmp_path, ARC)))
    assert r["arc"]["ki"] == [2]
    assert r["arc"]["sho"] == [3]
    assert r["turn_slide"] == 4
    assert r["arc"]["ketsu"] and r["arc"]["ketsu"][0] == 5


def test_the_lab_imperative_counts_as_a_turn(tmp_path):
    """"Keep the ladder, add one surface mode." is a proposal, not prose."""
    r = presentation_kishotenketsu_check(str(_deck(tmp_path, ARC)))
    assert r["checks"]["転 exists (a proposal after the complaint)"]


def test_a_title_band_without_a_placeholder_is_still_a_title(tmp_path):
    """The lab's builder draws its own band; `shapes.title` is None there."""
    r = presentation_kishotenketsu_check(str(_deck(tmp_path, ARC)))
    assert r["outline"][0]["slides"] == [2]
    # the title must not have leaked into the body and become the takeaway
    assert "The ladder" not in r["outline"][0]["takeaways"][0]


def test_hidden_slides_are_backup_whatever_they_are_called(tmp_path):
    """Nine hidden slides in the arc put the turn on the summary slide."""
    rows = ARC + [("Notation", ["Symbols."], "Every symbol on one slide."),
                  ("An L-shape", ["Non-tensor."], "The amplitude law holds to 1.5 %.")]
    r = presentation_kishotenketsu_check(
        str(_deck(tmp_path, rows, hidden_from=7, name="hidden.pptx")))
    assert r["backup_from_slide"] == 7
    assert r["content_slides"] == 5
    assert all(s < 7 for phase in r["arc"].values() for s in phase)


def test_the_turn_names_where_the_idea_came_from(tmp_path):
    r = presentation_kishotenketsu_check(str(_deck(tmp_path, ARC)))
    found = r["turn_inspiration"]
    assert found and found["slide"] == 4
    assert found["phrase"].lower() == "inspired by"
    assert found["credit"] == "[10]"
    assert r["checks"]["転 names where the idea came from (inspired-by, credited)"]


def test_an_uncredited_gesture_is_not_an_inspiration(tmp_path):
    """"as XFEM does" names a method, not a source: an acronym is no credit."""
    rows = [list(row) for row in ARC]
    rows[2][1] = ["We enrich the space, as XFEM does."]
    r = presentation_kishotenketsu_check(
        str(_deck(tmp_path, [tuple(x) for x in rows], name="vague.pptx")))
    assert r["turn_inspiration"] is None
    assert not r["checks"]["転 names where the idea came from (inspired-by, credited)"]
    assert any("どこから来たか" in c for c in r["comments"])


def test_an_inspiration_outside_the_turn_does_not_count(tmp_path):
    """Crediting the borrowed idea on the summary slide is too late."""
    rows = [list(row) for row in ARC]
    rows[2][1] = ["One space holds both."]
    rows[4][1] = ["Inspired by XFEM [10]."]
    r = presentation_kishotenketsu_check(
        str(_deck(tmp_path, [tuple(x) for x in rows], name="late.pptx")))
    assert r["turn_inspiration"] is None


def test_a_deck_too_short_to_have_an_arc(tmp_path):
    r = presentation_kishotenketsu_check(
        str(_deck(tmp_path, ARC[:2], name="short.pptx")))
    assert "error" in r


@pytest.mark.parametrize("hidden_from", [1, 2])
def test_backup_at_start_does_not_reintroduce_hidden_slides(tmp_path, hidden_from):
    result = presentation_kishotenketsu_check(
        str(_deck(tmp_path, ARC, hidden_from=hidden_from)))
    assert "only 0 content slides" in result["error"]


def test_bottom_banner_is_selected_independent_of_xml_order(tmp_path):
    path = _deck(tmp_path, ARC)
    prs = pptx.Presentation(path)
    for slide in prs.slides:
        top = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        top.text_frame.text = "Late-added body text"
    prs.save(path)
    result = presentation_kishotenketsu_check(str(path))
    assert result["turn_slide"] == 4
    assert result["assessment"] == "advisory-lexical-heuristic"
    assert any("English" in item for item in result["limitations"])


def test_recovered_tool_keeps_newer_title_style_api():
    import inspect
    from radia_mcp.presentation import tools
    assert callable(tools.presentation_kishotenketsu_check)
    assert "title_style" in inspect.signature(
        tools.presentation_check_slide_message_hierarchy).parameters


def _deck_with_furniture(tmp_path, rows, name="furniture.pptx"):
    """Build the lab layout with citations and page marks below the banner."""
    prs = pptx.Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for n, (title, body, takeaway) in enumerate(
            [("Cover", ["Talk"], "")] + rows, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        head = slide.shapes.add_textbox(Inches(0.6), Inches(0.1),
                                        Inches(12.0), Inches(0.7))
        head.text_frame.text = title
        for line in body:
            box = slide.shapes.add_textbox(Inches(0.6), Inches(1.5),
                                           Inches(12.0), Inches(3.0))
            box.text_frame.text = line
        if takeaway:
            band = slide.shapes.add_textbox(Inches(0.4), Inches(6.4),
                                            Inches(12.5), Inches(0.6))
            band.text_frame.paragraphs[0].add_run().text = takeaway
        cite = slide.shapes.add_textbox(Inches(0.4), Inches(7.05),
                                        Inches(9.0), Inches(0.3))
        cite.text_frame.text = "[1] Kameari et al., 2018"
        mark = slide.shapes.add_textbox(Inches(11.4), Inches(7.2),
                                        Inches(1.6), Inches(0.3))
        mark.text_frame.text = "Kindai University"
        page = slide.shapes.add_textbox(Inches(12.9), Inches(7.2),
                                        Inches(0.4), Inches(0.3))
        page.text_frame.text = str(n)
    out = tmp_path / name
    prs.save(str(out))
    return out


def test_page_furniture_is_not_read_as_the_takeaway(tmp_path):
    rows = list(ARC)
    rows.insert(2, ("Outline", ["So far the problem.",
                                "From here, the method."], ""))
    result = presentation_kishotenketsu_check(
        str(_deck_with_furniture(tmp_path, rows)))
    outline = next(item for item in result["outline"] if 4 in item["slides"])
    assert "4" not in outline["takeaways"]
    assert "Kindai University" not in outline["takeaways"]
    assert any("A circuit, exact at DC." in takeaway
               for phase in result["outline"]
               for takeaway in phase["takeaways"])


def test_no_x_does_y_is_a_complaint(tmp_path):
    rows = [
        ("The ladder", ["Kameari 2018 builds it."],
         "A circuit, DC exact: no rational model makes slope minus one half."),
        ("One space", ["Inspired by XFEM [10]: one matrix."],
         "Keep the ladder, add one surface mode."),
        ("A cylinder", ["Copper, 5 mm."],
         "Two unknowns hold it to 0.06 %."),
        ("Summary", ["Bulk plus surface."],
         "Keep the ladder, add the mode."),
    ]
    result = presentation_kishotenketsu_check(str(_deck(tmp_path, rows)))
    assert result["arc"]["sho"] == [2]
    assert result["turn_slide"] == 3


def test_named_backup_slide_is_reported_when_visible(tmp_path):
    path = _deck(tmp_path, [("Backup", ["Derivation"], "")],
                 name="visible-backup.pptx")
    result = presentation_check_qa_backup_slides(str(path), min_backup=1)
    assert result["named_but_visible"][0]["slide"] == 2
    assert result["warnings"]


def test_hidden_backup_slide_is_not_reported_as_visible(tmp_path):
    path = _deck(tmp_path, [("Backup", ["Derivation"], "")],
                 name="hidden-backup.pptx")
    prs = pptx.Presentation(path)
    prs.slides[1]._element.set("show", "0")
    prs.save(path)
    result = presentation_check_qa_backup_slides(str(path), min_backup=1)
    assert result["hidden_slides"] == 1
    assert result["named_but_visible"] == []
    assert result["warnings"] == []
