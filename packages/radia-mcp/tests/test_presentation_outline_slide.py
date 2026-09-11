"""Every major transition repeats the full agenda and highlights one item."""
import pytest

pptx = pytest.importorskip("pptx")

from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from radia_mcp.presentation._outline import (
    presentation_check_outline_slide,
)

ARC = [
    ("The ladder", ["Kameari 2018 builds it."], "A circuit, exact at DC."),
    ("Its far end", ["Ten rungs."], "A finite ladder cannot make that slope."),
    ("One space", ["Inspired by XFEM [10]."], "Keep the ladder, add one surface mode."),
    ("A cylinder", ["Copper, 5 mm."], "Two unknowns hold the band to 0.06 %."),
    ("Summary", ["Bulk plus surface."], "Keep the ladder, add the surface mode."),
]

AGENDA_ITEMS = ["1 Motivation", "2 Proposed method", "3 Results"]


def _agenda(active, title="Outline", items=None):
    return title, items or AGENDA_ITEMS, "", active


def _deck(tmp_path, rows, name="outline.pptx"):
    prs = pptx.Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for row in [("Cover", ["Talk"], "")] + rows:
        title, body, takeaway, *metadata = row
        active = metadata[0] if metadata else None
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        head = slide.shapes.add_textbox(Inches(0.6), Inches(0.1),
                                        Inches(12.0), Inches(0.7))
        head.text_frame.text = title
        box = slide.shapes.add_textbox(Inches(0.9), Inches(1.5),
                                       Inches(11.5), Inches(3.5))
        box.text_frame.clear()
        for index, line in enumerate(body):
            paragraph = (box.text_frame.paragraphs[0] if index == 0
                         else box.text_frame.add_paragraph())
            run = paragraph.add_run()
            run.text = line
            run.font.size = Pt(28)
            is_active = bool(active and active in line)
            run.font.bold = is_active
            run.font.color.rgb = (RGBColor(0xD6, 0x27, 0x28) if is_active
                                  else RGBColor(0x8A, 0x94, 0x9E))
        if takeaway:
            band = slide.shapes.add_textbox(Inches(0.4), Inches(6.4),
                                            Inches(12.5), Inches(0.6))
            band.text_frame.paragraphs[0].add_run().text = takeaway
    out = tmp_path / name
    prs.save(str(out))
    return out


def test_deck_without_section_dividers_fails(tmp_path):
    result = presentation_check_outline_slide(str(_deck(tmp_path, ARC)))
    assert result["score"] == 0.0
    assert result["section_dividers"] == []


def test_single_front_agenda_is_not_recurring_navigation(tmp_path):
    result = presentation_check_outline_slide(
        str(_deck(tmp_path, [_agenda("Motivation")] + ARC)))
    assert result["section_coverage"]["motivation"] == [2]
    assert result["section_coverage"]["method"] == []
    assert result["section_coverage"]["results"] == []
    assert result["score"] == 3.3


def test_three_section_agenda_is_below_recommended_range(tmp_path):
    rows = ([_agenda("Motivation")] + ARC[:2]
            + [_agenda("Proposed method")] + ARC[2:3]
            + [_agenda("Results")] + ARC[3:])
    result = presentation_check_outline_slide(str(_deck(tmp_path, rows)))
    assert result["score"] == 8.0
    assert result["section_coverage"] == {
        "motivation": [2], "method": [5], "results": [7]
    }


def test_sparse_section_cards_do_not_show_the_whole_route(tmp_path):
    rows = ([_agenda("Motivation", "1 Motivation", ["Motivation"])] + ARC[:2]
            + [_agenda("Proposed method", "2 Proposed method", ["Proposed method"])]
            + ARC[2:3] + [_agenda("Results", "3 Results", ["Results"])] + ARC[3:])
    result = presentation_check_outline_slide(str(_deck(tmp_path, rows)))
    assert result["score"] == 0.0


def test_full_agenda_without_unique_emphasis_fails(tmp_path):
    rows = [_agenda(None)] + ARC[:2] + [_agenda(None)] + ARC[2:3] + [_agenda(None)] + ARC[3:]
    result = presentation_check_outline_slide(str(_deck(tmp_path, rows)))
    assert result["score"] == 0.0
    assert all(not divider["highlighted_sections"]
               for divider in result["section_dividers"])


def test_section_names_are_examples_and_japanese_labels_work(tmp_path):
    items = ["1 背景と課題", "2 提案手法", "3 検証結果"]
    rows = ([_agenda("背景", items=items)] + ARC[:2]
            + [_agenda("提案", items=items)] + ARC[2:3]
            + [_agenda("結果", items=items)] + ARC[3:])
    result = presentation_check_outline_slide(str(_deck(tmp_path, rows)))
    assert result["score"] == 8.0


def test_four_section_custom_taxonomy_repeats_and_highlights_in_order(tmp_path):
    items = ["1 Motivation", "2 Theory", "3 Implementation", "4 Results"]
    rows = ([_agenda("Motivation", items=items)] + ARC[:1]
            + [_agenda("Theory", items=items)] + ARC[1:2]
            + [_agenda("Implementation", items=items)] + ARC[2:3]
            + [_agenda("Results", items=items)] + ARC[3:])
    result = presentation_check_outline_slide(str(_deck(tmp_path, rows)))
    assert result["score"] == 10.0
    assert result["scoring_mode"] == "repeated-agenda"
    assert result["recurring_agenda"]["items"] == items
    assert result["recurring_agenda"]["observed_highlight_order"] == items
    assert result["recurring_agenda"]["recommended_section_count"]
    assert result["recurring_agenda"]["concise_labels"]


def test_agenda_with_explanatory_taglines_does_not_receive_full_score(tmp_path):
    items = [
        "1 Motivation — why the ladder fails",
        "2 Proposed method — two spaces",
        "3 Results — four bodies",
        "4 Conclusion — use both",
    ]
    rows = ([_agenda("Motivation", items=items)] + ARC[:1]
            + [_agenda("Proposed method", items=items)] + ARC[1:2]
            + [_agenda("Results", items=items)] + ARC[2:4]
            + [_agenda("Conclusion", items=items)] + ARC[4:])
    result = presentation_check_outline_slide(str(_deck(tmp_path, rows)))
    assert result["score"] < 10.0
    assert not result["recurring_agenda"]["concise_labels"]
    assert not result["checks"]["Outline は説明文を付けず章名だけにする"]


def test_custom_section_titles_are_dividers_when_full_agenda_repeats(tmp_path):
    items = ["1 Context", "2 Theory", "3 Implementation", "4 Demonstration"]
    rows = ([_agenda("Context", title="1 Context", items=items)] + ARC[:1]
            + [_agenda("Theory", title="2 Theory", items=items)] + ARC[1:2]
            + [_agenda("Implementation", title="3 Implementation", items=items)]
            + ARC[2:3]
            + [_agenda("Demonstration", title="4 Demonstration", items=items)]
            + ARC[3:])
    result = presentation_check_outline_slide(str(_deck(tmp_path, rows)))
    assert result["score"] == 10.0
    assert result["scoring_mode"] == "repeated-agenda"
    assert [divider["slide"] for divider in result["section_dividers"]] == [2, 4, 6, 8]


def test_two_section_custom_taxonomy_is_supported(tmp_path):
    items = ["1 Theory", "2 Demonstration"]
    rows = ([_agenda("Theory", title="1 Theory", items=items)] + ARC[:2]
            + [_agenda("Demonstration", title="2 Demonstration", items=items)]
            + ARC[2:])
    result = presentation_check_outline_slide(str(_deck(tmp_path, rows)))
    assert result["score"] < 10.0
    assert result["recurring_agenda"]["observed_highlight_order"] == items


def test_custom_taxonomy_rejects_changed_order(tmp_path):
    items = ["1 Motivation", "2 Theory", "3 Implementation", "4 Results"]
    changed = ["1 Motivation", "3 Implementation", "2 Theory", "4 Results"]
    rows = ([_agenda("Motivation", items=items)] + ARC[:1]
            + [_agenda("Theory", items=items)] + ARC[1:2]
            + [_agenda("Implementation", items=changed)] + ARC[2:3]
            + [_agenda("Results", items=items)] + ARC[3:])
    result = presentation_check_outline_slide(str(_deck(tmp_path, rows)))
    assert result["score"] < 10.0
    assert not result["recurring_agenda"]["same_order"]


def test_deck_too_short_to_need_sections_says_so(tmp_path):
    result = presentation_check_outline_slide(str(_deck(tmp_path, ARC[:2])))
    assert "too few" in result["error"]
