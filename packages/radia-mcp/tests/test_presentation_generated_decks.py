"""Generated decks (2026-09-02): a slide that is one picture is flagged with
its effective dpi, and markdown bold markers left as text are found and
repaired into bold runs."""
from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from radia_mcp.presentation.tools import (
    presentation_apply_bold_markers,
    presentation_check_raster_slides,
    presentation_check_raw_math_markup,
)


def _picture(tmp_path: Path, px: tuple[int, int]) -> Path:
    path = tmp_path / f"pic_{px[0]}x{px[1]}.png"
    Image.new("RGB", px, "white").save(path)
    return path


def _deck(tmp_path: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    # slide 1: one full-slide picture, 1376 px across 13.333 in -> ~103 dpi
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    s1.shapes.add_picture(str(_picture(tmp_path, (1376, 768))), 0, 0,
                          width=prs.slide_width, height=prs.slide_height)
    # slide 2: a picture with text beside it -> editable, not flagged
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_picture(str(_picture(tmp_path, (1200, 800))), Inches(7), Inches(1),
                          width=Inches(6))
    box = s2.shapes.add_textbox(Inches(0.5), Inches(1), Inches(6), Inches(2))
    box.text_frame.text = "a real line of text"
    # slide 3: text with markdown bold left in it
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    box = s3.shapes.add_textbox(Inches(0.5), Inches(1), Inches(8), Inches(2))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "the crossover is **not fitted** here"
    run.font.size = Pt(24)
    run.font.name = "Times New Roman"
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def test_raster_slide_is_flagged_with_its_effective_dpi(tmp_path: Path):
    r = presentation_check_raster_slides(str(_deck(tmp_path)))
    assert r["n_slides"] == 3
    assert r["n_raster_slides"] == 1
    hit = r["raster_slides"][0]
    assert hit["slide"] == 1
    assert hit["pixels"] == [1376, 768]
    assert 100 < hit["effective_dpi"] < 106
    assert hit["soft_when_projected"] is True
    assert r["ok"] is False and r["editable"] is False
    assert "1 of 3" in r["verdict"]


def test_all_raster_deck_is_called_a_reference_not_a_deliverable(tmp_path: Path):
    prs = Presentation()
    for _ in range(2):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(str(_picture(tmp_path, (800, 600))), 0, 0,
                             width=prs.slide_width, height=prs.slide_height)
    path = tmp_path / "all_raster.pptx"
    prs.save(path)
    r = presentation_check_raster_slides(str(path))
    assert r["n_raster_slides"] == 2
    assert "not a deliverable" in r["verdict"]


def test_markdown_bold_is_reported_and_repaired(tmp_path: Path):
    deck = _deck(tmp_path)
    before = presentation_check_raw_math_markup(str(deck))
    kinds = {f["kind"] for f in before["findings"]}
    assert "markdown_emphasis" in kinds
    out = tmp_path / "repaired.pptx"
    r = presentation_apply_bold_markers(str(deck), str(out))
    assert r["ok"] is True
    assert r["bold_runs_made"] == 1
    assert r["leftover"] == []
    prs = Presentation(str(out))
    runs = [(rn.text, rn.font.bold, rn.font.size, rn.font.name)
            for p in prs.slides[2].shapes[0].text_frame.paragraphs for rn in p.runs]
    assert [t for t, *_ in runs] == ["the crossover is ", "not fitted", " here"]
    assert [b for _, b, *_ in runs] == [None, True, False]
    # size and typeface carried over to the new runs
    assert all(sz == Pt(24) and name == "Times New Roman" for _, _, sz, name in runs)
    after = presentation_check_raw_math_markup(str(out))
    assert "markdown_emphasis" not in {f["kind"] for f in after["findings"]}
