"""A negated claim only lands if the audience knows what is being denied.

Learned on the IGTE'26 deck (2026-09-03).  Its central slide was titled
"Schur complement, not a fit" and four slides said "nothing fitted": the
authors knew an earlier, unpublished construction that joined the bulk model
to the surface model with a fitted crossover, so to them "nothing is fitted"
was the whole point.  The audience had never seen that construction.  To
them "nothing fitted" denies a premise nobody had raised, and the value of
the claim is lost: they did not think anything needed fitting.

The rule: a negation ("no X", "nothing X-ed", "without X", "X-free", 「X
なし」「X 不要」「X はありません」) carries weight only when X was introduced
earlier in the deck as something one would otherwise do.  If X first appears
inside the negation itself, either introduce it first (one line: "the usual
repair fits a crossover frequency") or say what IS done instead ("the
projection decides the crossover").

:func:`presentation_check_negated_premise` walks the slides in show order,
collects negated claims from titles, bodies and notes, and reports those
whose premise has no earlier, non-negated mention.  The stemming is a
deliberate heuristic (documented in :func:`_stem`) so that "fitted",
"fitting" and "fit" count as one premise; a false positive costs one line of
reading, a miss costs a flat claim in front of the audience.
"""
from __future__ import annotations

import pathlib
import re

# English: the negation word, an optional auxiliary ("nothing IS fitted"),
# then the premise.  "not X" is not a claim pattern on purpose: "not" also
# negates the verbs of ordinary sentences ("does not replace").  It is still
# a negation for the other purpose, below: "not a fit" does not introduce
# fitting.
_AUX = r"(?:(?:is|are|was|were|be|been|being|get|gets|got|left|to|more|else)\s+)?"
_NEG_EN = re.compile(
    r"\b(?P<neg>no|nothing|without|never|zero)\s+" + _AUX +
    r"(?P<premise>[A-Za-z][A-Za-z-]{2,})",
    re.IGNORECASE,
)
_NOT_EN = re.compile(
    r"\bnot\s+(?:(?:a|an|the|any)\s+)?" + _AUX + r"(?P<premise>[A-Za-z][A-Za-z-]{2,})",
    re.IGNORECASE,
)
_FREE_EN = re.compile(r"\b(?P<premise>[A-Za-z]{3,})-(?P<neg>free)\b", re.IGNORECASE)
# Japanese: the premise is the run of kana/kanji right before the negation.
_NEG_JA = re.compile(
    r"(?P<premise>[ぁ-ゖァ-ヺ一-鿿ー]{2,8}?)"
    r"(?P<neg>なし|不要|はありません|がありません|はない|がない|は不要|を使わない|せず)"
)
# Words that follow a negation without being a premise anyone introduces.
_STOP = {
    "the", "this", "that", "these", "those", "one", "two", "three", "more",
    "other", "such", "any", "all", "each", "longer", "matter", "doubt",
    "less", "way", "need", "idea", "time", "part", "point", "reason",
    "problem", "question", "answer", "change", "wonder", "surprise",
}


def _stem(word: str) -> str:
    """fit / fits / fitted / fitting / fitter -> 'fit'; tune / tuned / tuning -> 'tun'.

    Strip a trailing s, then ed / ing / er, then a doubled final consonant
    left by the stripping, then a trailing e.  Not a stemmer for prose; a
    key that keeps the inflections of one technical verb together.
    """
    w = word.lower().strip("-")
    if w.endswith("s") and len(w) > 3:
        w = w[:-1]
    for suf in ("ing", "ed", "er"):
        if w.endswith(suf) and len(w) - len(suf) >= 3:
            w = w[: -len(suf)]
            break
    if len(w) >= 4 and w[-1] == w[-2] and w[-1] not in "aeiou":
        w = w[:-1]
    if w.endswith("e") and len(w) >= 4:
        w = w[:-1]
    return w


def _negations(text: str):
    """Yield (negation, premise, stem, span) for every negated claim in text."""
    for m in _NEG_EN.finditer(text):
        prem = m.group("premise")
        if prem.lower() in _STOP:
            continue
        yield m.group("neg"), prem, _stem(prem), m.span()
    for m in _FREE_EN.finditer(text):
        yield m.group("neg"), m.group("premise"), _stem(m.group("premise")), m.span()
    for m in _NEG_JA.finditer(text):
        prem = m.group("premise")
        yield m.group("neg"), prem, prem, m.span()


def _mentions(text: str, stem: str, premise: str, *, exclude: tuple | None = None) -> bool:
    """True if the premise appears in text outside a negation."""
    if re.search(r"[ぁ-鿿]", premise):
        idx = text.find(premise)
        while idx >= 0:
            tail = text[idx + len(premise): idx + len(premise) + 8]
            if not _NEG_JA.match(premise + tail):
                return True
            idx = text.find(premise, idx + 1)
        return False
    # a word inside any negation, the weak "not a fit" included, is not an
    # introduction of the premise
    negated = {span for _n, _p, _s, span in _negations(text)}
    negated |= {m.span() for m in _NOT_EN.finditer(text)}
    for m in re.finditer(r"[A-Za-z][A-Za-z-]{2,}", text):
        if _stem(m.group(0)) != stem:
            continue
        if any(a <= m.start() < b for a, b in negated):
            continue
        if exclude and exclude[0] <= m.start() < exclude[1]:
            continue
        return True
    return False


def check_negated_premises(slides: list[dict]) -> list[dict]:
    """Pure check on ``[{"slide": n, "title": str, "text": str}, ...]`` in order.

    A finding is a negated claim whose premise has no non-negated mention on
    any earlier slide, nor earlier on the same slide.
    """
    findings = []
    seen_text = ""
    for entry in slides:
        text = f"{entry.get('title', '')}\n{entry.get('text', '')}"
        for neg, prem, stem, span in _negations(text):
            before = text[: span[0]]
            if _mentions(seen_text, stem, prem) or _mentions(before, stem, prem):
                continue
            start = max(0, span[0] - 40)
            findings.append({
                "slide": entry.get("slide"),
                "title": entry.get("title", ""),
                "negation": neg,
                "premise": prem,
                "context": text[start: span[1] + 40].replace("\n", " ").strip(),
            })
        seen_text += "\n" + text
    return findings


def _slide_entries(prs, include_notes: bool):
    for i, slide in enumerate(prs.slides, 1):
        try:
            hidden = slide._element.get("show") == "0"
        except Exception:
            hidden = False
        if hidden:
            continue
        title = ""
        parts = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            txt = shape.text_frame.text
            if not title and getattr(shape, "is_placeholder", False):
                try:
                    if shape.placeholder_format.type is not None and "TITLE" in str(shape.placeholder_format.type):
                        title = txt.strip()
                        continue
                except Exception:
                    pass
            parts.append(txt)
        if not title and parts:
            title = parts[0].strip().splitlines()[0][:80] if parts[0].strip() else ""
        if include_notes and slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
        yield {"slide": i, "title": title, "text": "\n".join(parts)}


def presentation_check_negated_premise(pptx_path: str,
                                       include_notes: bool = True,
                                       max_findings: int = 40) -> dict:
    """否定形の主張のうち、否定している前提を聴衆がまだ知らないものを検出する。

    「no X」「nothing X-ed」「without X」「X-free」「X なし」「X 不要」の類は、
    X を「普通ならやること」として聴衆が既に知っているときだけ効く。X が
    その否定文の中で初めて現れるなら、聴衆には「そもそも X が要るとは思って
    いなかった」ので主張の価値が伝わらない（IGTE'26 の "nothing fitted" が
    これ: 当てはめで接続した未発表の先行構成を聴衆は知らない）。

    表示されるスライドを順に読み、題・本文・（既定で）ノートの否定形の主張の
    うち、前提が先行スライドにも同じスライドの前段にも否定なしで現れないもの
    を返す。英語の語幹処理は fitted / fitting / fit を一つに数える程度の発見
    的なもの。

    Args:
        pptx_path: 点検する .pptx。
        include_notes: ノートも読む（既定 True）。
        max_findings: 返す指摘の上限。

    Returns: ``{"ok", "n_findings", "findings", "hint"}``。``findings`` は
    slide / title / negation / premise / context。
    """
    try:
        import pptx as _pptx
    except ImportError:
        return {"error": "python-pptx not installed."}
    path = pathlib.Path(pptx_path)
    if not path.exists():
        return {"error": f"file not found: {pptx_path}"}
    prs = _pptx.Presentation(str(path))
    findings = check_negated_premises(list(_slide_entries(prs, include_notes)))
    return {
        "ok": not findings,
        "n_findings": len(findings),
        "findings": findings[:max_findings],
        "hint": (
            "前提を先に一行で立てる（「通常の補修は接続周波数を当てはめる」）か、"
            "否定をやめて代わりに何をするかを言う（「接続は射影が決める」）。"
        ),
    }
