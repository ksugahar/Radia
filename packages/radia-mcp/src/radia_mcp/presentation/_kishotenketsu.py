"""Does the talk turn? A 起承転結 check that reads the bottom line, not the title.

This supplementary check targets decks with a short topic title and a bottom
takeaway. Message-style titles remain supported by the separate hierarchy
checker; this tool does not impose one title style on all presentations.

This reads the takeaway line instead, where the lab's format puts the assertion,
and uses vocabulary cues to suggest four moves:

  起  the object and the existing approach, stated without complaint
  承  the complaint: what the existing approach cannot do
  転  the turn: a proposal that answers that complaint
  結  evidence for the proposal, its limits, and the summary

The turn is also where the audience most wants to know WHERE THE IDEA CAME
FROM. A proposal that arrives from nowhere has to be taken on trust; a
proposal that says "the extended finite element method solves the ordinary
basis and the enrichment in one matrix, and that is what we borrowed" is
followed on the first hearing, because the listener already believes the
borrowed idea. Sugahara asked for this in the IGTE'26 deck on 2026-09-06:
「inspired ロジックは伝わりやすい」. So the turn is also checked for a named,
credited outside idea.

These are advisory lexical cues, not semantic proof: matching words or a
percentage does not demonstrate an answer or valid evidence. English word
overlap is not suitable for scoring Japanese-only prose. Review the slide
content before accepting any suggested diagnosis.
"""
from __future__ import annotations

import re

# The complaint: something the current approach cannot do, or does badly.
_COMPLAINT = re.compile(
    r"(?i)\bcannot\b|\bcan not\b|\bfails?\b|\bno[t]? \w+ed\b|\bstops? at\b"
    r"|\bno\b[^.;!?]{0,48}?\b(?:makes?|produces?|gives?|reach(?:es)?|holds?"
    r"|covers?|works?|does)\b|\bneither\b"
    r"|\bby hand\b|\bceiling\b|\blimit(?:ation)?s?\b|\bprice\b|\bwrong\b"
    r"|\bmisses\b|\bbreaks?\b|\bnot the\b|できない|限界|欠陥|課題|問題")
# The turn: an instruction or a resolution, often answering the complaint.
# The imperative may open ANY sentence of the takeaway, not only the
# first: "One basis cannot make that slope. Use two, one projected on the
# other." states the lack and the proposal in one line, and an anchored ^
# missed it -- the turn then fell through to the summary slide.
_TURN = re.compile(
    r"(?i)(?:^|(?<=[.;!?] ))(?:do not|use|take|add|replace|instead|keep"
    r"|mix|combine|join|enrich|project|let|drop|solve)\b|\buse both\b"
    r"|\bcomes out of\b|\bremoves?\b|\bnothing is (?:left to be )?fit"
    r"|\bnothing is fitted\b|\binstead\b|\bwe propose\b|提案|代わりに")
# Evidence: a takeaway that carries a measured number.
_NUMBER = re.compile(r"\d+(?:\.\d+)?\s*(?:%|x\b|×|倍)")
# Where the idea came from: the turn names an outside idea it borrowed.
_INSPIRED = re.compile(
    r"(?i)\binspired by\b|\bborrow(?:ed|ing|s)?\b|\bin the spirit of\b"
    r"|\bafter the \w+ of\b|\badapted from\b|\bfollowing \w+'s\b"
    r"|\bthe way \w+ does\b|\bas \w+ does\b|\bthis came from\b"
    r"|\bcame from\b|着想|示唆を得|に倣|に学|由来|ヒント")
# ...and credits it, so the borrowing is attributable, not a vague gesture.
# An acronym is NOT a credit: "as XFEM does" names a method, not a source.
_CREDIT = re.compile(r"\[\d{1,3}\]|\bet al\.|[A-Z][a-z]{2,}(?:'s)?\s+\d{4}"
                     r"|先生|らの|ら\s*\[")

# Text below this fraction is page furniture, not a takeaway: citation line,
# institution mark, and page number. The lab takeaway band sits above it.
FOOTER_FROM = 0.92


def _moves(takeaway: str) -> set[str]:
    t = (takeaway or "").strip()
    out = set()
    if _COMPLAINT.search(t):
        out.add("complaint")
    if _TURN.search(t):
        out.add("turn")
    if _NUMBER.search(t):
        out.add("evidence")
    return out


def _overlap(a: str, b: str) -> float:
    """Content-word overlap, as a fraction of the shorter side."""
    stop = {"the", "a", "an", "is", "are", "of", "to", "and", "in", "on",
            "that", "it", "this", "with", "for", "not", "no", "one", "at",
            "from", "by", "its", "has", "have", "be", "so", "but", "which"}
    wa = {w for w in re.findall(r"[a-z]+", a.lower()) if w not in stop}
    wb = {w for w in re.findall(r"[a-z]+", b.lower()) if w not in stop}
    if not wa or not wb:
        return 0.0
    return len(wa & wb) / min(len(wa), len(wb))


def read_deck(pptx_path: str, backup_title: str = "Backup"):
    """Read slide titles, takeaways, text, and the visible-talk boundary.

    The outline and arc checks share this reader so drawn title bands, hidden
    backup slides, and page furniture have one definition.
    """
    from pptx import Presentation

    prs = Presentation(pptx_path)
    slide_h = float(prs.slide_height or 0)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        eligible = [
            shape for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
            and not (shape.top is not None and slide_h
                     and int(shape.top) > FOOTER_FROM * slide_h)
        ]
        if slide.shapes.title is not None:
            title = slide.shapes.title.text.strip()
        else:
            heads = [
                (int(shape.top), shape.text_frame.text.strip())
                for shape in eligible
                if shape.top is not None
                and int(shape.top) < 0.25 * slide_h
                and "\n" not in shape.text_frame.text.strip()
            ]
            title = min(heads)[1] if heads else ""

        texts = [shape.text_frame.text.strip() for shape in eligible]
        body = [text for text in texts if text != title]
        # XML insertion order is not visual order; select the lowest eligible
        # text box while excluding the footer strip above.
        banners = [
            (int(shape.top), int(shape.left), shape.text_frame.text.strip())
            for shape in eligible
            if shape.text_frame.text.strip() != title
        ]
        takeaway = max(banners)[2].splitlines()[-1].strip() if banners else ""
        hidden = slide._element.get("show") == "0"
        slides.append({"slide": i, "title": title, "takeaway": takeaway,
                       "text": "\n".join(body), "hidden": hidden})

    cut = next((index for index, slide in enumerate(slides)
                if slide["hidden"] or slide["title"].strip() == backup_title),
               len(slides))
    main = slides[1:cut]  # drop the title card; never reintroduce backup
    return slides, cut, main


def presentation_kishotenketsu_check(pptx_path: str,
                                     backup_title: str = "Backup") -> dict:
    """Suggest a 起承転結 arc using lexical cues, not semantic validation.

    backup_title : slides from this one onward are treated as backup and left
                   out; they are answers to questions, not part of the arc.
    """
    try:
        from pptx import Presentation
    except ImportError:
        return {"error": "python-pptx not installed."}

    slides, cut, main = read_deck(pptx_path, backup_title)
    if len(main) < 4:
        return {"error": f"only {len(main)} content slides; too few to have an arc."}

    for s in main:
        s["moves"] = sorted(_moves(s["takeaway"]))

    first_complaint = next((k for k, s in enumerate(main)
                            if "complaint" in s["moves"]), None)
    turn = None
    if first_complaint is not None:
        turn = next((k for k, s in enumerate(main)
                     if k > first_complaint and "turn" in s["moves"]), None)
    evidence = [k for k, s in enumerate(main)
                if "evidence" in s["moves"] and (turn is None or k > turn)]

    # 起 is whatever precedes the first complaint; 結 begins at the first
    # evidence slide after the turn.
    arc = {}
    if first_complaint is not None:
        arc["ki"] = [main[k]["slide"] for k in range(first_complaint)]
        end_sho = turn if turn is not None else len(main)
        arc["sho"] = [main[k]["slide"] for k in range(first_complaint, end_sho)]
    if turn is not None:
        end_ten = evidence[0] if evidence else len(main)
        arc["ten"] = [main[k]["slide"] for k in range(turn, end_ten)]
        arc["ketsu"] = [main[k]["slide"] for k in range(end_ten, len(main))]

    # does the turn answer the complaint, and does the close echo the turn?
    # A deck may raise several complaints and answer them on different slides.
    # Pairing only the first of each reports "the turn answers nothing" for a
    # talk that answers everything, just not in that pair.
    answers, unanswered = 0.0, []
    complaints = [k for k, sl in enumerate(main) if "complaint" in sl["moves"]]
    proposals = [k for k, sl in enumerate(main)
                 if "turn" in sl["moves"] and (turn is None or k >= turn)]
    for c in complaints:
        if turn is not None and c > turn:
            continue                      # a limitation stated after the turn
        best = max((_overlap(main[c]["takeaway"], main[t]["takeaway"])
                    for t in proposals), default=0.0)
        answers = max(answers, best)
        if best < 0.15:
            unanswered.append({"slide": main[c]["slide"],
                               "takeaway": main[c]["takeaway"],
                               "best_overlap": round(best, 2)})
    echo = 0.0
    if turn is not None:
        echo = max((_overlap(main[t]["takeaway"], main[-1]["takeaway"])
                    for t in proposals), default=0.0)

    # 転 が「どこから来たか」を名乗っているか。借り物の考えを credit つきで
    # 名指しすると、聞き手は既に信じている物に相乗りできるので初聴で通る。
    inspiration = None
    for k in (arc.get("ten") or []):
        slide = next(s for s in main if s["slide"] == k)
        phrase = _INSPIRED.search(slide["text"] or "")
        credit = _CREDIT.search(slide["text"] or "")
        if phrase and credit:
            inspiration = {"slide": k, "phrase": phrase.group(0),
                           "credit": credit.group(0)}
            break

    checks = {
        "起 exists (setup before any complaint)":
            first_complaint is not None and first_complaint > 0,
        "承 exists (a stated complaint)": first_complaint is not None,
        "転 exists (a proposal after the complaint)": turn is not None,
        "結 exists (measured evidence after the turn)": bool(evidence),
        "転 answers 承 (shared content words)": answers >= 0.15,
        "結 echoes 転 (the close returns to the proposal)": echo >= 0.15,
        "転 lands in the middle (not the last quarter)":
            turn is not None and turn < len(main) * 0.75,
        "転 names where the idea came from (inspired-by, credited)":
            inspiration is not None,
    }
    score = round(10.0 * sum(checks.values()) / len(checks), 1)

    comments = [f"{'OK  ' if v else 'FAIL'} {k}" for k, v in checks.items()]
    if first_complaint is None:
        comments.append(
            "承 が無い。従来法の何ができないかを下端文で言い切ると、転が効く。")
    elif turn is None:
        comments.append(
            "転が見つからない。提案を下端文に命令形か『〜が決まる』の形で置く。")
    for u in unanswered:
        comments.append(
            f"スライド {u['slide']} の欠陥に答える提案が見当たらない"
            f"（最大の語重なり {u['best_overlap']:.2f}）: 「{u['takeaway'][:52]}」")
    if turn is not None and echo < 0.15:
        comments.append(
            f"まとめが転に戻っていない（語の重なり {echo:.2f}）。"
            "最後の下端文は提案の言い直しにする。")
    if turn is not None and inspiration is None:
        comments.append(
            "転が「どこから来たか」を言っていない。借りた外の考えを "
            "inspired by / 着想を得た の形で名指しし、[n] などの出典を"
            "同じ枚に添えると、提案は初聴で通る。聞き手は既に信じている"
            "考えに相乗りできるので、新奇さを一から売り込まずに済む。")
    elif inspiration is not None:
        comments.append(
            f"OK   転 {inspiration['slide']} 枚目が出所を名乗っている"
            f"（「{inspiration['phrase']}」/ {inspiration['credit']}）。")

    outline = [
        {"phase": ph, "slides": arc.get(ph, []),
         "takeaways": [s["takeaway"] for s in main if s["slide"] in arc.get(ph, [])]}
        for ph in ("ki", "sho", "ten", "ketsu")
    ]
    return {
        "assessment": "advisory-lexical-heuristic",
        "limitations": [
            "Cue matches and word overlap do not establish scientific validity or semantic coherence.",
            "Overlap scoring uses English words; Japanese-only decks require human review.",
            "The first slide is treated as a cover and the lowest text box as the takeaway.",
        ],
        "score": score,
        "score_max": 10,
        "content_slides": len(main),
        "backup_from_slide": slides[cut]["slide"] if cut < len(slides) else None,
        "arc": {k: v for k, v in arc.items()},
        "turn_slide": main[turn]["slide"] if turn is not None else None,
        "turn_position": round((turn + 1) / len(main), 2) if turn is not None else None,
        "turn_inspiration": inspiration,
        "turn_answers_complaint": round(answers, 2),
        "unanswered_complaints": unanswered,
        "close_echoes_turn": round(echo, 2),
        "checks": checks,
        "comments": comments,
        "outline": outline,
        "hint": "転の inspired-by だけは下端文ではなく枚全体の文字を見る"
                "（出所は本文に書く）。ほかはタイトルではなく下端 takeaway "
                "の語句を補助点検する。タイトル形式は変更しない。"
                "得点だけで発表の論理や品質を判定しない。",
        "source": "起承転結（四段構成）を学会発表に適用。"
                  "presentation_check_slide_message_hierarchy と対で使う: "
                  "あちらは 1 枚の中の役割分担、こちらは 枚をまたぐ筋。",
    }
