"""Decks that arrive as pictures, and markdown that arrived as text.

Two things learned from a Gemini Notebook deck on 2026-09-02, when it was
put beside the lab's hand-built deck of the same talk:

* A generated deck is often one full-slide raster per slide: nothing on it
  can be edited, and at 1376 px across a 45 cm slide it projects at ~78 dpi.
  Its value is as a *reference for how to show things* -- font size as a
  fraction of the slide, where the takeaway sits, labels with arrows instead
  of legends -- not as a deliverable. :func:`presentation_check_raster_slides`
  says which slides are such pictures and at what effective dpi, so nobody
  presents one by mistake and nobody tries to edit one.

* A deck built from markdown may keep ``**bold**`` as literal asterisks when
  the converter knows headings, bullets and math but not emphasis; the raw
  markup check reports it as ``markdown_emphasis`` and
  :func:`presentation_apply_bold_markers` turns the markers into bold runs.
"""
from __future__ import annotations

import pathlib
import re
from copy import deepcopy

_BOLD = re.compile(r"\*\*([^*\n]{1,120})\*\*")


def presentation_check_raster_slides(pptx_path: str,
                                     cover_fraction: float = 0.85,
                                     min_effective_dpi: float = 150.0) -> dict:
    """Find slides that are a single picture with no editable text.

    A slide counts as a raster slide when its pictures cover at least
    ``cover_fraction`` of the slide area and it carries no text run. For each,
    the picture's pixel size is turned into an effective dpi at the size it is
    shown; below ``min_effective_dpi`` it will look soft when projected.

    Such a deck is a reference for how to show things, not a deliverable:
    nothing on it can be edited, corrected or re-fonted. Say so before anyone
    tries.

    Args:
        pptx_path: the .pptx to inspect.
        cover_fraction: picture area over slide area from which a slide is
            called a raster slide.
        min_effective_dpi: floor for a picture shown at its placed size.

    Returns: ``{"ok", "n_slides", "n_raster_slides", "raster_slides",
    "editable", "verdict"}``; each raster slide lists its pixel size, the
    displayed width in cm and the effective dpi.
    """
    try:
        import pptx as _pptx
    except ImportError:
        return {"error": "python-pptx not installed."}
    path = pathlib.Path(pptx_path)
    if not path.exists():
        return {"error": f"file not found: {pptx_path}"}
    try:
        from PIL import Image
    except ImportError:  # pragma: no cover - PIL ships with matplotlib
        Image = None

    prs = _pptx.Presentation(str(path))
    slide_area = float(prs.slide_width) * float(prs.slide_height)
    raster = []
    for no, slide in enumerate(prs.slides, 1):
        pics, text_runs = [], 0
        for sh in slide.shapes:
            if sh.shape_type == 13:  # PICTURE
                pics.append(sh)
            elif sh.has_text_frame:
                text_runs += sum(1 for p in sh.text_frame.paragraphs
                                 for r in p.runs if r.text.strip())
        if not pics or text_runs:
            continue
        covered = sum(float(p.width) * float(p.height) for p in pics) / slide_area
        if covered < cover_fraction:
            continue
        big = max(pics, key=lambda p: float(p.width) * float(p.height))
        width_cm = float(big.width) / 360000.0
        px_w = px_h = None
        if Image is not None:
            try:
                import io
                with Image.open(io.BytesIO(big.image.blob)) as im:
                    px_w, px_h = im.size
            except Exception:  # noqa: BLE001 - an unreadable blob is reported as such
                pass
        dpi = (px_w / (float(big.width) / 914400.0)) if px_w else None
        raster.append({
            "slide": no,
            "shape": big.name,
            "covered_fraction": round(covered, 3),
            "pixels": [px_w, px_h] if px_w else None,
            "displayed_width_cm": round(width_cm, 2),
            "effective_dpi": round(dpi, 1) if dpi else None,
            "soft_when_projected": (dpi is not None and dpi < min_effective_dpi),
        })

    n = len(prs.slides._sldIdLst)
    n_r = len(raster)
    if n_r == 0:
        verdict = "every slide carries editable text; nothing to flag"
    elif n_r == n:
        verdict = ("every slide is one picture: a reference for how to show things, "
                   "not a deliverable -- nothing on it can be edited or re-fonted")
    else:
        verdict = f"{n_r} of {n} slides are pictures with no editable text"
    return {
        "ok": n_r == 0,
        "n_slides": n,
        "n_raster_slides": n_r,
        "raster_slides": raster,
        "editable": n_r == 0,
        "verdict": verdict,
        "hint": ("Measure a raster deck instead of copying it: text height in px "
                 "over px/cm gives its font size, to compare at equal slide width."),
    }


def apply_bold_markers_to_shape(shape) -> int:
    """Turn ``**text**`` inside a shape's runs into bold runs. Returns how many."""
    from pptx.text.text import _Run

    made = 0
    for par in shape.text_frame.paragraphs:
        for run in list(par.runs):
            text = run.text
            if "**" not in text:
                continue
            parts = _BOLD.split(text)  # plain, bold, plain, bold, ...
            run.text = parts[0]
            prev = run._r
            for k, seg in enumerate(parts[1:], start=1):
                if not seg:
                    continue
                el = deepcopy(run._r)
                prev.addnext(el)
                prev = el
                new = _Run(el, run._parent)
                new.text = seg
                new.font.bold = (k % 2 == 1)
                made += k % 2
    return made


def presentation_apply_bold_markers(pptx_path: str, out_path: str = "") -> dict:
    """Turn literal ``**bold**`` markers in slide text into bold runs.

    Each marked run is split into plain / bold / plain ... runs; the new runs
    are copies of the original, so size, colour and typeface carry over and
    only bold changes. Writes to ``out_path`` (default: in place) and reports
    any marker that could not be converted, which is worse than no emphasis.

    Args:
        pptx_path: the .pptx to repair.
        out_path: where to write; empty means overwrite ``pptx_path``.

    Returns: ``{"ok", "bold_runs_made", "leftover", "written"}``.
    """
    try:
        import pptx as _pptx
    except ImportError:
        return {"error": "python-pptx not installed."}
    path = pathlib.Path(pptx_path)
    if not path.exists():
        return {"error": f"file not found: {pptx_path}"}
    prs = _pptx.Presentation(str(path))
    made = 0
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                made += apply_bold_markers_to_shape(sh)
    leftover = [
        {"slide": no, "shape": sh.name, "text": p.text[:80]}
        for no, slide in enumerate(prs.slides, 1)
        for sh in slide.shapes if sh.has_text_frame
        for p in sh.text_frame.paragraphs if "**" in p.text
    ]
    target = pathlib.Path(out_path) if out_path else path
    prs.save(str(target))
    return {
        "ok": not leftover,
        "bold_runs_made": made,
        "leftover": leftover,
        "written": str(target),
    }
