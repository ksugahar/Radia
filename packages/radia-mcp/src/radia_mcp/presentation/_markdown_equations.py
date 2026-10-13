"""Script-first PPTX assembly; all mathematical serialization belongs to Eqnedit64."""
from __future__ import annotations

from copy import deepcopy
import importlib
import os
from pathlib import Path
import re
import tempfile
import zipfile


MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
A14_NS = "http://schemas.microsoft.com/office/drawing/2010/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
MIN_EQNEDIT64 = (3, 1, 4)


def _backend():
    try:
        backend = importlib.import_module("eqnedit64")
    except ImportError as exc:
        raise RuntimeError("Install eqnedit64>=3.1.4 with its native PowerPoint math backend") from exc
    version = getattr(backend, "__version__", "")
    if not re.fullmatch(r"\d+\.\d+\.\d+", version) or tuple(map(int, version.split("."))) < MIN_EQNEDIT64:
        raise RuntimeError(f"eqnedit64>=3.1.4 is required; loaded {version!r}")
    if not callable(getattr(backend, "export_powerpoint_math", None)):
        raise RuntimeError("Installed Eqnedit64 has no native PowerPoint math export")
    return backend


def _blocks(markdown: str):
    """Parse the intentionally small, explicit script-first block contract."""
    if not isinstance(markdown, str) or not markdown.strip() or len(markdown) > 1_000_000:
        raise ValueError("markdown must be non-empty and at most one million characters")
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        index += 1
        if not line.strip():
            continue
        if line.startswith("```"):
            if not re.fullmatch(r"```[\w+-]*", line):
                raise ValueError("Unsupported code fence")
            code = []
            while index < len(lines) and lines[index] != "```":
                code.append(lines[index]); index += 1
            if index == len(lines):
                raise ValueError("Unclosed code fence")
            index += 1
            yield "code", "\n".join(code), 0
            continue
        heading = re.fullmatch(r"#{1,6} (.+)", line)
        if heading:
            if "$" in heading[1] or "`" in heading[1]:
                raise ValueError("Slide headings must be plain text")
            yield "heading", heading[1], 0
            continue
        if line.startswith(">"):
            notes = [line[1:].lstrip()]
            while index < len(lines) and lines[index].startswith(">"):
                notes.append(lines[index][1:].lstrip()); index += 1
            yield "notes", "\n".join(notes), 0
            continue
        bullet = re.fullmatch(r"[-+*] (.+)", line)
        number = re.fullmatch(r"(\d+)\. (.+)", line)
        if bullet:
            yield "bullet", bullet[1], 0
            continue
        if number:
            if not 1 <= int(number[1]) <= 32767:
                raise ValueError("Numbered-list item is out of range")
            yield "numbered", number[2], int(number[1])
            continue
        if line.startswith((" ", "\t", "|", "~", "#")) or re.fullmatch(r"[-*_]{3,}", line):
            raise ValueError("Nested lists, tables, thematic breaks and indented blocks are unsupported")
        paragraph = [line]
        # Ordinary paragraph continuations, including multiline display TeX.
        while index < len(lines) and lines[index].strip():
            next_line = lines[index]
            if next_line.startswith(("#", ">", "```")) or re.match(r"(?:[-+*] |\d+\. )", next_line):
                break
            if next_line.startswith((" ", "\t", "|", "~")) and sum(s.count("$$") for s in paragraph) % 2 == 0:
                raise ValueError("Unsupported paragraph continuation")
            paragraph.append(next_line); index += 1
        yield "paragraph", "\n".join(paragraph), 0


def _pieces(text: str):
    """Yield text/inline/display segments; never reinterpret code as TeX."""
    plain = []
    index = 0
    while index < len(text):
        char = text[index]
        if char == "\\" and index + 1 < len(text) and text[index + 1] in "$`":
            plain.append(text[index + 1]); index += 2; continue
        if char == "`":
            end = text.find("`", index + 1)
            if end < 0 or end == index + 1:
                raise ValueError("Unclosed or unsupported inline code span")
            if plain:
                value = "".join(plain)
                if re.search(r"!\[|\[[^\]]*\]\(|\*\*|__|<[A-Za-z!/]", value):
                    raise ValueError("Links, images, emphasis and raw HTML are unsupported")
                yield "text", value
                plain = []
            yield "literal", text[index + 1:end]
            index = end + 1; continue
        if char != "$":
            plain.append(char); index += 1; continue
        if plain:
            value = "".join(plain)
            if re.search(r"!\[|\[[^\]]*\]\(|\*\*|__|<[A-Za-z!/]", value):
                raise ValueError("Links, images, emphasis and raw HTML are unsupported")
            yield "text", value
            plain = []
        delimiter = "$$" if text.startswith("$$", index) else "$"
        end = index + len(delimiter)
        while True:
            end = text.find(delimiter, end)
            if end < 0:
                raise ValueError("Unclosed math delimiter; escape literal currency dollars")
            before = end - 1
            while before >= 0 and text[before] == "\\":
                before -= 1
            if (end - before - 1) % 2 == 0:
                break
            end += len(delimiter)
        tex = text[index + len(delimiter):end]
        if not tex.strip() or (delimiter == "$" and ("\n" in tex or tex != tex.strip())):
            raise ValueError("Empty or malformed math delimiter")
        yield "display" if delimiter == "$$" else "inline", tex
        index = end + len(delimiter)
    if plain:
        value = "".join(plain)
        if re.search(r"!\[|\[[^\]]*\]\(|\*\*|__|<[A-Za-z!/]", value):
            raise ValueError("Links, images, emphasis and raw HTML are unsupported")
        yield "text", value


def presentation_markdown_to_pptx(markdown: str, output_path: str,
                                  title: str = "", font_pt: float = 20.0,
                                  timeout_s: float = 30.0) -> dict:
    """Build an editable equation deck with Eqnedit64>=3.1.4.

    Supported subset: plain headings (slides), plain paragraphs, flat bullets
    and numbered lists, fenced/inline code, > script lines (speaker notes),
    $TeX$ inline and $$TeX$$ display math in slide body and notes. Escape
    currency as \\$. Links, images, emphasis, tables, nested lists, raw HTML and
    malformed/unsupported TeX raise. Native math stays at the 18 pt Office
    contract; font_pt controls prose. Modern PowerPoint is required: math
    shapes use a14 with no image/text fallback. No clipboard is touched.
    """
    blocks = list(_blocks(markdown))
    if not isinstance(title, str) or "$" in title:
        raise ValueError("title must be plain text")
    if not 8 <= float(font_pt) <= 72 or not 0 < float(timeout_s) <= 120:
        raise ValueError("font_pt must be 8..72 and timeout_s in (0, 120]")
    backend = _backend()
    try:
        from lxml import etree
        from pptx import Presentation
        from pptx.util import Pt
    except ImportError as exc:
        raise RuntimeError("Install radia-mcp[document,equation] for PPTX assembly") from exc
    output = Path(output_path).expanduser().resolve()
    if output.suffix.lower() != ".pptx":
        raise ValueError("output_path must end with .pptx")
    output.parent.mkdir(parents=True, exist_ok=True)
    counts = {"body_math": 0, "notes_math": 0}
    with tempfile.TemporaryDirectory(prefix="radia-pptx-", dir=output.parent) as scratch:
        cache = {}
        prs = Presentation()
        if title:
            cover = prs.slides.add_slide(prs.slide_layouts[0])
            cover.shapes.title.text = title
            cover.placeholders[1].text = ""
        slide = None
        states = {}

        def new_slide(heading):
            nonlocal slide
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = heading
            slide.placeholders[1].text_frame.clear()

        def paragraph(frame, kind="paragraph", number=0):
            key = frame._txBody
            used = states.get(key, False)
            result = frame.add_paragraph() if used else frame.paragraphs[0]
            states[key] = True
            props = result._p.get_or_add_pPr()
            for child in list(props):
                if child.tag in {f"{{{DRAWING_NS}}}buNone", f"{{{DRAWING_NS}}}buChar", f"{{{DRAWING_NS}}}buAutoNum"}:
                    props.remove(child)
            if kind == "bullet":
                bullet = etree.SubElement(props, f"{{{DRAWING_NS}}}buChar")
                bullet.set("char", "•")
            elif kind == "numbered":
                bullet = etree.SubElement(props, f"{{{DRAWING_NS}}}buAutoNum")
                bullet.set("type", "arabicPeriod"); bullet.set("startAt", str(number))
            else:
                etree.SubElement(props, f"{{{DRAWING_NS}}}buNone")
            # Bullet properties precede defRPr in the DrawingML schema.
            bullet = props[-1]
            props.remove(bullet)
            before = next((i for i, node in enumerate(props)
                           if node.tag in {f"{{{DRAWING_NS}}}tabLst", f"{{{DRAWING_NS}}}defRPr", f"{{{DRAWING_NS}}}extLst"}), len(props))
            props.insert(before, bullet)
            result.font.size = Pt(float(font_pt))
            return result

        def write(frame, text, kind, number, count_key):
            current = None
            segments = [("text", text)] if kind == "code" else list(_pieces(text))
            for mode, value in segments:
                if current is None or mode == "display":
                    current = paragraph(frame, kind, number)
                    if mode == "display":
                        current.space_before = Pt(8)
                        current.space_after = Pt(8)
                if mode in {"text", "literal"}:
                    run = current.add_run(); run.text = value
                    if kind == "code" or mode == "literal":
                        run.font.name = "Consolas"; run.font.size = Pt(14)
                else:
                    if value not in cache:
                        path = Path(scratch) / f"math-{len(cache)}.pptxmath"
                        backend.export_powerpoint_math(value, path, timeout_s=float(timeout_s))
                        parser = etree.XMLParser(resolve_entities=False, no_network=True)
                        root = etree.fromstring(path.read_bytes(), parser)
                        if root.tag != f"{{{MATH_NS}}}oMathPara" or any(
                            not element.tag.startswith((f"{{{MATH_NS}}}", f"{{{DRAWING_NS}}}"))
                            or element.tag == f"{{{MATH_NS}}}t" for element in root.iter()
                        ):
                            raise RuntimeError("Eqnedit64 returned an invalid PowerPoint math dialect")
                        cache[value] = root
                    root = deepcopy(cache[value])
                    payload = root if mode == "display" else root.find(f"{{{MATH_NS}}}oMath")
                    if payload is None:
                        raise RuntimeError("Eqnedit64 output has no Office math")
                    wrapper = etree.SubElement(current._p, f"{{{A14_NS}}}m", nsmap={"a14": A14_NS})
                    wrapper.append(payload)
                    counts[count_key] += 1
                    if mode == "display":
                        current = None

        for kind, text, number in blocks:
            if kind == "heading":
                new_slide(text); continue
            if slide is None:
                new_slide(title)
            if kind == "notes":
                notes = slide.notes_slide
                # Office does not consistently inherit notes placeholder geometry
                # once a math shape is enclosed in AlternateContent. Materialize
                # the master's geometry for all cloned notes placeholders.
                for placeholder in notes.placeholders:
                    master = next((p for p in prs.notes_master.placeholders
                                   if p.placeholder_format.type == placeholder.placeholder_format.type), None)
                    if master is None:
                        raise RuntimeError("Notes placeholder has no matching master")
                    placeholder._element.replace(placeholder._element.spPr,
                                                 deepcopy(master._element.spPr))
                    if not placeholder.has_text_frame and master.has_text_frame:
                        placeholder._element.append(deepcopy(master.text_frame._txBody))
                    if placeholder.has_text_frame and master.has_text_frame:
                        target = placeholder.text_frame._txBody
                        target.replace(target.bodyPr, deepcopy(master.text_frame._txBody.bodyPr))
                frame = notes.notes_text_frame
                if frame._txBody not in states:
                    frame.clear()
                write(frame, text, "paragraph", 0, "notes_math")
            else:
                write(slide.placeholders[1].text_frame, text, kind, number, "body_math")
        # Math is a DrawingML extension, so the compatibility envelope belongs
        # around its containing shape, not inside a paragraph. Latest only:
        # omit the optional Fallback instead of inventing an image substitute.
        for part in list(prs.part.package.iter_parts()):
            element = getattr(part, "_element", None)
            if element is None:
                continue
            for shape in list(element.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}sp")):
                if next(shape.iter(f"{{{A14_NS}}}m"), None) is None:
                    continue
                parent = shape.getparent(); index = parent.index(shape)
                envelope = etree.Element(f"{{{MC_NS}}}AlternateContent", nsmap={"mc": MC_NS, "a14": A14_NS})
                choice = etree.SubElement(envelope, f"{{{MC_NS}}}Choice", Requires="a14")
                choice.append(shape); parent.insert(index, envelope)
        candidate = Path(scratch) / "deck.pptx"
        prs.save(candidate)
        reopened = Presentation(candidate)
        if len(reopened.slides) != len(prs.slides):
            raise RuntimeError("PPTX reopen changed slide count")
        with zipfile.ZipFile(candidate) as archive:
            actual = 0
            for name in archive.namelist():
                if re.fullmatch(r"ppt/(?:slides/slide|notesSlides/notesSlide)\d+\.xml", name):
                    root = etree.fromstring(archive.read(name))
                    actual += len(list(root.iter(f"{{{MATH_NS}}}oMath")))
            if actual != sum(counts.values()):
                raise RuntimeError("Saved PPTX lost editable equations")
        os.replace(candidate, output)
    return {"ok": True, "output_path": str(output), "slides": len(prs.slides),
            **counts, "backend": "Eqnedit64", "backend_version": backend.__version__,
            "compatibility": "modern PowerPoint with a14 math; no fallback"}
